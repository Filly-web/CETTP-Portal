import os
import re
import io
import sqlite3
import tempfile
from datetime import datetime, date
from pathlib import Path

import qrcode
from flask import (
    Blueprint, render_template, request, redirect, url_for, flash, current_app,
    send_file, send_from_directory, abort
)
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_DIR = BASE_DIR / "templates" / "efd_portal"
DATA_DIR = Path(os.environ.get("EFD_PORTAL_DATA_DIR", os.environ.get("PORTAL_DATA_DIR", str(BASE_DIR / "instance"))))
DATA_DIR.mkdir(parents=True, exist_ok=True)
DATABASE = DATA_DIR / "portal_history.db"
LOGO_FILE = BASE_DIR / "britam-logo.png"
ALLOWED_EXTENSIONS = {"docx"}
MAX_CONTENT_LENGTH = 15 * 1024 * 1024  # 15 MB

efd_bp = Blueprint("efd_portal", __name__, template_folder=str(TEMPLATE_DIR))


def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS processing_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                process_date TEXT NOT NULL,
                filename TEXT NOT NULL,
                verification_code TEXT NOT NULL,
                receipt_time TEXT NOT NULL
            )
        """)


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def normalize_time(value):
    value = (value or "").strip()
    if re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d:[0-5]\d", value):
        return value
    if re.fullmatch(r"(?:[01]\d|2[0-3])[0-5]\d[0-5]\d", value):
        return f"{value[0:2]}:{value[2:4]}:{value[4:6]}"
    raise ValueError("Enter receipt time as HH:MM:SS or HHMMSS.")


def make_verification_url(code, receipt_time):
    clean_time = receipt_time.replace(":", "")
    return f"https://verify.tra.go.tz/{code}_{clean_time}"


def make_qr_image(url):
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=6,
        border=2,
    )
    qr.add_data(url)
    qr.make(fit=True)
    image = qr.make_image(fill_color="black", back_color="white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _style_borderless_table(table):
    """Remove visible table borders so the details and QR look like one clean block."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders")
    if borders is None:
        from docx.oxml import OxmlElement
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _fill_efd_table(table, code, receipt_time, verification_url, qr_buffer):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Inches, Pt
    table.autofit = False
    table.columns[0].width = Inches(4.6)
    table.columns[1].width = Inches(0.8)
    left, right = table.rows[0].cells
    left.width = Inches(4.6)
    right.width = Inches(0.8)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading = p.add_run("EFD Verification Details")
    heading.bold = True
    heading.font.size = Pt(10)

    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run(f"TRA Verification Code: {code}").font.size = Pt(9)
    p.add_run(f"\nReceipt Time: {receipt_time}").font.size = Pt(9)
    p.add_run(f"\nVerification URL: {verification_url}").font.size = Pt(8)

    p_qr = right.paragraphs[0]
    p_qr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_qr.paragraph_format.space_after = Pt(0)
    qr_buffer.seek(0)
    p_qr.add_run().add_picture(qr_buffer, width=Inches(0.55))
    _style_borderless_table(table)


def add_efd_section(document, code, receipt_time, verification_url, qr_buffer):
    """
    Replaces <<EFD_SECTION>> with a compact, borderless two-column block:
    verification details on the left and a smaller QR code immediately beside them.
    If no marker is found, the block is appended to the end of the document.
    """
    marker = "<<EFD_SECTION>>"

    # Search ordinary document paragraphs and insert the block directly after the marker.
    for paragraph in document.paragraphs:
        if marker in paragraph.text:
            for run in paragraph.runs:
                run.text = ""
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            table = document.add_table(rows=1, cols=2)
            _fill_efd_table(table, code, receipt_time, verification_url, qr_buffer)
            paragraph._p.addnext(table._tbl)
            return

    # Search table cells. Use a nested two-column table in the same cell.
    for outer_table in document.tables:
        for row in outer_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if marker in paragraph.text:
                        for run in paragraph.runs:
                            run.text = ""
                        table = cell.add_table(rows=1, cols=2)
                        _fill_efd_table(table, code, receipt_time, verification_url, qr_buffer)
                        return

    # No marker: append the compact block at the end of the document.
    document.add_paragraph()
    table = document.add_table(rows=1, cols=2)
    _fill_efd_table(table, code, receipt_time, verification_url, qr_buffer)


@efd_bp.route("/britam-logo.png")
def britam_logo():
    if not LOGO_FILE.is_file():
        abort(404)
    return send_from_directory(str(BASE_DIR), "britam-logo.png")


@efd_bp.route("/", methods=["GET"])
def dashboard():
    with get_db() as conn:
        total = conn.execute(
            "SELECT COUNT(*) AS n FROM processing_history"
        ).fetchone()["n"]
        today = conn.execute(
            "SELECT COUNT(*) AS n FROM processing_history WHERE process_date = ?",
            (date.today().isoformat(),)
        ).fetchone()["n"]
        recent = conn.execute(
            "SELECT created_at, filename, verification_code "
            "FROM processing_history ORDER BY id DESC LIMIT 8"
        ).fetchall()

    return render_template("efd_portal/index.html", total=total, today=today, recent=recent)


@efd_bp.route("/process", methods=["POST"])
def process_invoice():
    uploaded = request.files.get("invoice")
    verification_code = (request.form.get("verification_code") or "").strip()
    raw_time = request.form.get("verification_time", "")

    if not uploaded or not uploaded.filename:
        flash("Please select a Word invoice (.docx).", "error")
        return redirect(url_for("efd_portal.dashboard"))

    if not allowed_file(uploaded.filename):
        flash("Only Microsoft Word .docx files are supported.", "error")
        return redirect(url_for("efd_portal.dashboard"))

    if not verification_code:
        flash("Enter the TRA verification code.", "error")
        return redirect(url_for("efd_portal.dashboard"))

    # Avoid path separators and unsafe names in downloads/history.
    original_filename = secure_filename(uploaded.filename) or "invoice.docx"

    try:
        receipt_time = normalize_time(raw_time)
        verification_url = make_verification_url(verification_code, receipt_time)

        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / original_filename
            output_name = f"{Path(original_filename).stem}_processed.docx"
            output_path = Path(temp_dir) / output_name
            uploaded.save(input_path)

            try:
                document = Document(str(input_path))
            except Exception as exc:
                raise ValueError("The uploaded file could not be opened as a valid .docx document.") from exc

            qr_buffer = make_qr_image(verification_url)
            add_efd_section(
                document, verification_code, receipt_time,
                verification_url, qr_buffer
            )
            document.save(str(output_path))

            with get_db() as conn:
                conn.execute("""
                    INSERT INTO processing_history
                    (created_at, process_date, filename, verification_code, receipt_time)
                    VALUES (?, ?, ?, ?, ?)
                """, (
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    date.today().isoformat(),
                    original_filename,
                    verification_code,
                    receipt_time,
                ))

            return send_file(
                output_path,
                as_attachment=True,
                download_name=output_name,
                mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

    except ValueError as exc:
        flash(str(exc), "error")
        return redirect(url_for("efd_portal.dashboard"))
    except Exception:
        current_app.logger.exception("Invoice processing failed")
        flash("Invoice processing failed. Check that the file is a valid .docx and try again.", "error")
        return redirect(url_for("efd_portal.dashboard"))


@efd_bp.route("/history", methods=["GET"])
def history():
    with get_db() as conn:
        rows = conn.execute("""
            SELECT id, created_at, filename, verification_code, receipt_time
            FROM processing_history ORDER BY id DESC LIMIT 250
        """).fetchall()
    return render_template("efd_portal/history.html", rows=rows)


@efd_bp.app_errorhandler(413)
def file_too_large(_error):
    flash("The uploaded file is too large. Maximum size is 15 MB.", "error")
    return redirect(url_for("efd_portal.dashboard"))


init_db()
